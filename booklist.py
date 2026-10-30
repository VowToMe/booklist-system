# -*- coding: utf-8 -*-
"""书单底本维护脚本（唯一数据源 = booklist.xlsx；CSV 为可 grep 的镜像）。

用法：
  python booklist.py init                        # 建底本（已存在则不动）
  python booklist.py add books.json [--sheet X]  # 追加（按书名/名称去重）
  python booklist.py list [--sheet X] [--status 待读] [--cat 法学]
  python booklist.py stats
  python booklist.py export
选项：
  --master <path>   覆盖底本路径（自检用，不动真实底本）
"""
import argparse
import csv
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

MASTER = Path(__file__).resolve().parent / "booklist.xlsx"

BOOK_COLS = ["序号", "书名", "作者", "类别", "状态", "优先级", "来源", "加入日期", "为什么加", "备注"]
ALT_COLS = ["序号", "名称", "类型", "平台", "状态", "优先级", "来源", "加入日期", "为什么加", "备注"]

SHEETS = {
    "书单": {
        "cols": BOOK_COLS,
        "name_key": "书名",
        "widths": {"A": 6, "B": 34, "C": 18, "D": 12, "E": 8, "F": 8, "G": 22, "H": 12, "I": 46, "J": 30},
        "dv": {"D": ["法学", "社科", "历史", "文学", "哲学", "心理", "商业/经济", "科学", "方法/工具", "英文原版", "其他"],
               "E": ["待读", "在读", "读完", "弃读", "存疑"],
               "F": ["高", "中", "低"]},
        "defaults": {"类别": "其他", "状态": "待读", "优先级": "中"},
    },
    "其他资源": {
        "cols": ALT_COLS,
        "name_key": "名称",
        "widths": {"A": 6, "B": 30, "C": 12, "D": 16, "E": 8, "F": 8, "G": 22, "H": 12, "I": 46, "J": 30},
        "dv": {"C": ["视频频道", "课程", "播客", "文章", "工具", "其他"],
               "E": ["未看", "在看", "看完", "弃"],
               "F": ["高", "中", "低"]},
        "defaults": {"类型": "其他", "状态": "未看", "优先级": "中"},
    },
}


def norm_title(t) -> str:
    """去书名号、标点、空白、大小写，用于判重。"""
    s = str(t or "")
    s = re.sub(r"[《》〈〉「」『』\"'“”‘’\s:：·,，.。!！?？\-—_()（）\[\]【】]", "", s)
    return s.lower()


def today() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def backup(path: Path) -> Path:
    """改动/删除已有行之前留一份快照。add 只追加，不备份。"""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dst = path.with_name(f"{path.stem}.bak-{stamp}{path.suffix}")
    shutil.copy2(path, dst)
    return dst


def load(path: Path, sheet: str):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return wb, ws, []
    header = [str(c) if c is not None else "" for c in rows[0]]
    items = []
    for r in rows[1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        items.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
    return wb, ws, items


def init(path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if path.exists():
        return "已存在，未改动"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for sheet, cfg in SHEETS.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = sheet
        first = False
        ws.append(cfg["cols"])
        for i in range(1, len(cfg["cols"]) + 1):
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="3F4A57")
            c.alignment = Alignment(vertical="center")
        for col, w in cfg["widths"].items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{chr(64 + len(cfg['cols']))}1"
        for col, options in cfg["dv"].items():
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}5000")
    wb.save(path)
    return "已建：" + " / ".join(SHEETS)


def _row_of(item, cfg, idx):
    vals = {"序号": idx}
    for col in cfg["cols"][1:]:
        v = item.get(col)
        if v in (None, ""):
            v = cfg["defaults"].get(col, "")
        if col == "加入日期" and not v:
            v = today()
        vals[col] = v
    return [vals[c] for c in cfg["cols"]]


def set_filter(ws, ncols):
    """让自动筛选覆盖到数据最后一行（只筛表头的话 Excel 里筛选是坏的）。"""
    last = chr(64 + ncols)
    ws.auto_filter.ref = f"A1:{last}{max(ws.max_row, 1)}"


def export_csv(path: Path, sheet: str, items, cols):
    base = path.with_suffix("")
    csv_path = base.parent / (base.name + (".csv" if sheet == "书单" else f"-{sheet}.csv"))
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for it in items:
            w.writerow(["" if it.get(c) is None else it.get(c) for c in cols])
    return csv_path


def add(path: Path, sheet: str, payload):
    cfg = SHEETS[sheet]
    key = cfg["name_key"]
    init(path)
    wb, ws, items = load(path, sheet)
    have = {norm_title(b.get(key)) for b in items}
    added, skipped, bad = [], [], []
    for raw in payload:
        if isinstance(raw, str):
            item = {key: raw}
        else:
            item = dict(raw)
        name = str(item.get(key, "")).strip()
        if not name:
            bad.append(item)
            continue
        if norm_title(name) in have:
            skipped.append(name)
            continue
        have.add(norm_title(name))
        ws.append(_row_of(item, cfg, len(items) + len(added) + 1))
        added.append(name)
    if added:
        set_filter(ws, len(cfg["cols"]))
    wb.save(path)
    _, _, items = load(path, sheet)
    export_csv(path, sheet, items, cfg["cols"])
    print(json.dumps({"表": sheet, "新增": added, "重复跳过": skipped, "无效": len(bad), "总数": len(items)},
                     ensure_ascii=False))


def pad(s, width):
    s = str(s if s not in (None, "") else "-")
    w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)
    return s + " " * max(0, width - w)


def show(path: Path, sheet: str, status=None, cat=None, level=None):
    cfg = SHEETS[sheet]
    key = cfg["name_key"]
    _, _, items = load(path, sheet)
    rows = [b for b in items
            if (not status or b.get(cfg["cols"][4]) == status)
            and (not cat or b.get(cfg["cols"][3]) == cat)
            and (not level or b.get(cfg["cols"][5]) == level)]
    print(pad(key, 36) + pad(cfg["cols"][2], 18) + pad(cfg["cols"][3], 12) + pad(cfg["cols"][4], 8) + pad("优先级", 8))
    for b in rows:
        print(pad(b.get(key), 36) + pad(b.get(cfg["cols"][2]), 18) + pad(b.get(cfg["cols"][3]), 12)
              + pad(b.get(cfg["cols"][4]), 8) + pad(b.get(cfg["cols"][5]), 8))
    print(f"-- {len(rows)} / {len(items)} 条")


def remove_rows(path: Path, sheet: str, names):
    """按名称删除行并重排序号（只删指定行，不动其余数据）。"""
    cfg = SHEETS[sheet]
    key = cfg["name_key"]
    targets = {norm_title(n) for n in names}
    wb, ws, items = load(path, sheet)
    removed = [b.get(key) for b in items if norm_title(b.get(key)) in targets]
    keep = backup(path)
    for r in range(ws.max_row, 1, -1):
        if norm_title(ws.cell(row=r, column=2).value) in targets:
            ws.delete_rows(r)
    for i, r in enumerate(range(2, ws.max_row + 1), start=1):
        ws.cell(row=r, column=1, value=i)
    set_filter(ws, len(cfg["cols"]))
    wb.save(path)
    _, _, items = load(path, sheet)
    export_csv(path, sheet, items, cfg["cols"])
    print(json.dumps({"表": sheet, "删除": removed, "未找到": [n for n in names if norm_title(n) not in
                     {norm_title(x) for x in removed}], "剩余": len(items), "备份": keep.name},
                     ensure_ascii=False))


def clear_sheet(path: Path, sheet: str):
    """清空一张 sheet 的数据行（保留表头与格式）。"""
    cfg = SHEETS[sheet]
    wb, ws, items = load(path, sheet)
    n = len(items)
    keep = backup(path)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    set_filter(ws, len(cfg["cols"]))
    wb.save(path)
    export_csv(path, sheet, [], cfg["cols"])
    print(json.dumps({"表": sheet, "清空": n, "剩余": 0, "备份": keep.name}, ensure_ascii=False))


def update_rows(path: Path, sheet: str, names, changes):
    """改已有行的字段（状态 / 优先级 / 类别 / 备注 / 作者 / 为什么加）。

    按书名认领（去书名号、标点、大小写的归一化比对），写前自动备份。
    这是「管辖书单」的写入端：`add` 只能追加，改已有行只能走这里。
    """
    cfg = SHEETS[sheet]
    key = cfg["name_key"]
    cols = cfg["cols"]
    changes = {k: v for k, v in (changes or {}).items() if v not in (None, "")}
    if not changes:
        raise SystemExit("set 至少要给一个要改的字段（--status/--level/--cat/--note/--author/--why）")
    allowed = set(cols) - {"序号", key}
    bad = sorted(set(changes) - allowed)
    if bad:
        raise SystemExit(f"不能改这几列：{bad}；可改：{sorted(allowed)}")
    for letter, vals in (cfg.get("dv") or {}).items():
        cname = cols[ord(letter) - 65]
        v = changes.get(cname)
        if v and v not in vals:
            raise SystemExit(f"「{cname}」只能是 {'/'.join(vals)}，收到「{v}」")
    wb, ws, items = load(path, sheet)
    pos = {norm_title(b.get(key)): r for r, b in enumerate(items, start=2)}
    done, missed = [], []
    for n in names:
        r = pos.get(norm_title(n))
        if r is None:
            missed.append(n)
            continue
        diff = {}
        for c, v in changes.items():
            old = ws.cell(row=r, column=cols.index(c) + 1).value
            if str(old or "") != str(v):
                diff[c] = [old or "", v]
        if not diff:
            done.append({"书名": ws.cell(row=r, column=2).value, "变更": "无（已是该值）"})
            continue
        for c, (old, new) in diff.items():
            ws.cell(row=r, column=cols.index(c) + 1, value=new)
        done.append({"书名": ws.cell(row=r, column=2).value, "变更": diff})
    keep = None
    if any(isinstance(d.get("变更"), dict) for d in done):
        keep = backup(path)
        set_filter(ws, len(cols))
        wb.save(path)
        _, _, items = load(path, sheet)
        export_csv(path, sheet, items, cfg["cols"])
    print(json.dumps({"表": sheet, "改了": done, "没找到": missed, "总数": len(items),
                      "备份": keep.name if keep else None}, ensure_ascii=False))


def stats(path: Path):
    out = {}
    for sheet, cfg in SHEETS.items():
        _, _, items = load(path, sheet)
        if not items:
            continue

        def tally(key):
            d = {}
            for b in items:
                k = str(b.get(key) or "-")
                d[k] = d.get(k, 0) + 1
            return d
        out[sheet] = {"总数": len(items), "状态": tally(cfg["cols"][4]),
                      cfg["cols"][3]: tally(cfg["cols"][3])}
    print(json.dumps(out, ensure_ascii=False))


def fix(path: Path):
    """只规范化筛选区/序号，不动任何内容。"""
    wb = None
    for sheet, cfg in SHEETS.items():
        wb, ws, items = load(path, sheet)
        for i, r in enumerate(range(2, ws.max_row + 1), start=1):
            ws.cell(row=r, column=1, value=i)
        set_filter(ws, len(cfg["cols"]))
        wb.save(path)
    print("已规范化：" + " / ".join(SHEETS))


def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass
    p = argparse.ArgumentParser()
    p.add_argument("cmd", choices=["init", "add", "list", "stats", "export", "remove", "clear", "fix", "set"])
    p.add_argument("input", nargs="?")
    p.add_argument("--master", default=str(MASTER))
    p.add_argument("--sheet", default="书单", choices=list(SHEETS))
    p.add_argument("--status")
    p.add_argument("--cat")
    p.add_argument("--level")
    p.add_argument("--note")
    p.add_argument("--author")
    p.add_argument("--why")
    p.add_argument("--name", action="append", default=[])
    a = p.parse_args()
    path = Path(a.master)
    if path.parent.exists() and list(path.parent.glob("~$" + path.name)):
        print("底本正被 Excel/WPS 打开，先关掉再跑。")
        return 2
    if a.cmd == "init":
        print(init(path))
    elif a.cmd == "add":
        add(path, a.sheet, json.loads(Path(a.input).read_text(encoding="utf-8")))
    elif a.cmd == "list":
        show(path, a.sheet, a.status, a.cat, a.level)
    elif a.cmd == "stats":
        stats(path)
    elif a.cmd == "export":
        for sheet, cfg in SHEETS.items():
            _, _, items = load(path, sheet)
            print(export_csv(path, sheet, items, cfg["cols"]))
    elif a.cmd == "remove":
        remove_rows(path, a.sheet, a.name)
    elif a.cmd == "clear":
        clear_sheet(path, a.sheet)
    elif a.cmd == "fix":
        fix(path)
    elif a.cmd == "set":
        cols = SHEETS[a.sheet]["cols"]
        update_rows(path, a.sheet, a.name,
                    {c: v for c, v in ((cols[4], a.status), (cols[5], a.level),
                                       (cols[3], a.cat), (cols[9], a.note),
                                       (cols[2], a.author), (cols[8], a.why)) if v})
    return 0


if __name__ == "__main__":
    sys.exit(main())
